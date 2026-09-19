"""Resume ingestion: text extraction, OCR detection and structured field pull.

Design rules:
  * Never raise out of ``parse_resume`` — every failure becomes a structured
    error on the returned object so one bad file cannot lose a whole batch.
  * Never invent content. If a section is absent, the field stays empty and the
    UI renders "Not mentioned in resume".
"""

from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional

from .config import NOT_MENTIONED
from .utils import (
    content_hash,
    dewrap_lines,
    estimate_years_from_dates,
    extract_candidate_name,
    extract_total_years,
    normalize_text,
    split_evidence_units,
    strip_bullets,
    text_fingerprint,
)

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}

# A text layer shorter than this on a PDF almost always means a scanned image.
MIN_CHARS_PER_PAGE_FOR_TEXT_PDF = 80
MIN_TOTAL_CHARS = 120


class ResumeParseError(Exception):
    """Raised internally and converted into a structured status."""

    def __init__(self, message: str, status: str = "failed"):
        super().__init__(message)
        self.status = status


@dataclass
class ParsedResume:
    """Everything extracted from one resume file."""

    file_name: str
    status: str = "processed"          # processed | failed | ocr_required | duplicate
    error: str = ""
    text: str = ""
    page_count: int = 0
    char_count: int = 0
    content_sha: str = ""
    text_sha: str = ""
    extraction_method: str = ""
    candidate_name: str = ""
    name_extracted: bool = False
    email: str = ""
    phone: str = ""
    location: str = ""
    notice_period: str = ""
    total_years: Optional[float] = None
    years_source: str = ""
    job_titles: List[str] = field(default_factory=list)
    companies: List[str] = field(default_factory=list)
    skills: List[str] = field(default_factory=list)
    education: List[str] = field(default_factory=list)
    certifications: List[str] = field(default_factory=list)
    responsibilities: List[str] = field(default_factory=list)
    sections: Dict[str, str] = field(default_factory=dict)
    evidence_units: List[str] = field(default_factory=list)
    duplicate_of: str = ""

    @property
    def ok(self) -> bool:
        return self.status == "processed"

    def summary_line(self) -> str:
        bits = [f"{self.page_count} page(s)", f"{self.char_count} chars", self.extraction_method]
        return " · ".join(b for b in bits if b)


# ---------------------------------------------------------------------------
# Raw text extraction
# ---------------------------------------------------------------------------

def _extract_pdf(data: bytes) -> tuple[str, int, str]:
    """Extract text from a PDF using every available engine, best result wins.

    Both engines are tried whenever the first one comes back thin. A PDF is only
    called "scanned" once every engine has failed to find a text layer —
    otherwise one engine returning nothing (which happens on some builds and
    some PDF producers) would send a perfectly readable document to OCR.
    """
    attempts: list[tuple[str, int, str]] = []   # (text, pages, method)
    last_error: Optional[Exception] = None

    def substantive(text: str) -> bool:
        return len(normalize_text(text)) >= MIN_TOTAL_CHARS

    # --- Engine 1: PyMuPDF ---
    try:
        import pymupdf  # type: ignore
    except ImportError:  # pragma: no cover - older wheels expose `fitz` only
        try:
            import fitz as pymupdf  # type: ignore
        except ImportError:
            pymupdf = None  # type: ignore

    if pymupdf is not None:
        try:
            with pymupdf.open(stream=data, filetype="pdf") as doc:
                if getattr(doc, "needs_pass", False):
                    # Try the common "owner password only" case before failing.
                    if not doc.authenticate(""):
                        raise ResumeParseError(
                            "Password-protected PDF — cannot be read without the password.")
                pages = [page.get_text("text") or "" for page in doc]
                text = "\n".join(pages)
                if substantive(text):
                    return text, len(pages), "pymupdf"
                attempts.append((text, len(pages), "pymupdf"))
        except ResumeParseError:
            raise
        except Exception as exc:  # noqa: BLE001 - try the next engine
            last_error = exc
            logger.debug("PyMuPDF extraction failed: %s", exc)

    # --- Engine 2: pdfplumber ---
    try:
        import pdfplumber  # type: ignore

        with pdfplumber.open(io.BytesIO(data)) as pdf:
            pages = [(page.extract_text() or "") for page in pdf.pages]
            text = "\n".join(pages)
            if substantive(text):
                return text, len(pages), "pdfplumber"
            attempts.append((text, len(pages), "pdfplumber"))
    except Exception as exc:  # noqa: BLE001
        last_error = exc
        logger.debug("pdfplumber extraction failed: %s", exc)

    if attempts:
        # Every engine ran but none found much. Hand back the best of them and
        # let the caller decide whether this is a scan or simply a sparse file.
        best = max(attempts, key=lambda a: len(normalize_text(a[0])))
        return best

    message = str(last_error) if last_error else "unknown error"
    if "password" in message.lower() or "encrypt" in message.lower():
        raise ResumeParseError("Password-protected PDF — cannot be read without the password.")
    raise ResumeParseError(f"Corrupted or unreadable PDF ({message[:160]}).")


def _ocr_pdf(data: bytes) -> tuple[str, str]:
    """OCR a scanned PDF when the tooling is available.

    Returns ``(text, method)``. Raises ResumeParseError with status
    ``ocr_required`` when OCR cannot run — we flag the file rather than guess at
    its contents.
    """
    try:
        import pytesseract  # type: ignore
        from PIL import Image  # type: ignore
    except ImportError:
        raise ResumeParseError(
            "Scanned/image-based PDF. OCR libraries (pytesseract, Pillow) are not installed.",
            status="ocr_required",
        )

    try:
        import pymupdf  # type: ignore
    except ImportError:  # pragma: no cover
        try:
            import fitz as pymupdf  # type: ignore
        except ImportError:
            raise ResumeParseError(
                "Scanned/image-based PDF. PyMuPDF is required to rasterise pages for OCR.",
                status="ocr_required",
            )

    try:
        pytesseract.get_tesseract_version()
    except Exception:  # noqa: BLE001
        raise ResumeParseError(
            "No text layer found, so this looks like a scanned/image-based PDF, and OCR is "
            "unavailable (the Tesseract binary is not installed). If this file does show "
            "selectable text when you open it, re-save or re-export it as a PDF and try "
            "again; otherwise install Tesseract to enable OCR.",
            status="ocr_required",
        )

    try:
        chunks: List[str] = []
        with pymupdf.open(stream=data, filetype="pdf") as doc:
            for page in doc:
                pix = page.get_pixmap(dpi=220)
                image = Image.open(io.BytesIO(pix.tobytes("png")))
                chunks.append(pytesseract.image_to_string(image) or "")
        text = "\n".join(chunks)
    except Exception as exc:  # noqa: BLE001
        raise ResumeParseError(f"OCR failed: {str(exc)[:160]}", status="ocr_required")

    if len(normalize_text(text)) < MIN_TOTAL_CHARS:
        raise ResumeParseError(
            "Scanned/image-based PDF. OCR ran but produced too little text to evaluate.",
            status="ocr_required",
        )
    return text, "ocr(tesseract)"


def _extract_docx(data: bytes) -> tuple[str, int, str]:
    try:
        import docx  # type: ignore

        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ResumeParseError(f"Corrupted or unreadable DOCX ({str(exc)[:160]}).")

    parts: List[str] = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts), 0, "python-docx"


def _extract_txt(data: bytes) -> tuple[str, int, str]:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding), 0, f"text/{encoding}"
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace"), 0, "text/replace"


def extract_text(file_name: str, data: bytes) -> tuple[str, int, str]:
    """Dispatch on extension and return ``(text, page_count, method)``."""
    if not data:
        raise ResumeParseError("Empty file (0 bytes).")

    suffix = Path(file_name).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ResumeParseError(
            f"Unsupported format '{suffix or 'unknown'}'. Supported: PDF, DOCX, TXT.")

    if suffix == ".pdf":
        text, pages, method = _extract_pdf(data)
        clean = normalize_text(text)
        per_page = len(clean) / max(pages, 1)
        if len(clean) < MIN_TOTAL_CHARS or per_page < MIN_CHARS_PER_PAGE_FOR_TEXT_PDF:
            # Almost no text layer -> scanned document. Try OCR, else flag it.
            ocr_text, ocr_method = _ocr_pdf(data)
            return ocr_text, pages, ocr_method
        return text, pages, method

    if suffix == ".docx":
        return _extract_docx(data)
    return _extract_txt(data)


# ---------------------------------------------------------------------------
# Structured field extraction
# ---------------------------------------------------------------------------

SECTION_PATTERNS: Dict[str, re.Pattern] = {
    "summary": re.compile(r"^\s*(professional\s+summary|summary|profile|about\s+me|objective|career\s+objective)\s*:?\s*$", re.I),
    "experience": re.compile(r"^\s*(work\s+experience|professional\s+experience|experience|employment(\s+history)?|career\s+history|work\s+history)\s*:?\s*$", re.I),
    "education": re.compile(r"^\s*(education|academic(s|\s+background|\s+qualifications)?|qualifications)\s*:?\s*$", re.I),
    "skills": re.compile(r"^\s*(skills|technical\s+skills|core\s+skills|key\s+skills|skills\s*&\s*tools|competencies|core\s+competencies|tools?(\s*&\s*technologies)?|technologies)\s*:?\s*$", re.I),
    "certifications": re.compile(r"^\s*(certifications?|licenses?\s*(&|and)?\s*certifications?|courses|training)\s*:?\s*$", re.I),
    "projects": re.compile(r"^\s*(projects?|key\s+projects|selected\s+projects)\s*:?\s*$", re.I),
    "achievements": re.compile(r"^\s*(achievements?|accomplishments?|awards?|highlights)\s*:?\s*$", re.I),
}

_EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
_PHONE_RE = re.compile(r"(?:(?:\+|00)\d{1,3}[\s.-]?)?(?:\(?\d{3,5}\)?[\s.-]?)\d{3}[\s.-]?\d{3,4}\b")
_LOCATION_RE = re.compile(
    r"(?:^|\n|\||·|•)\s*(?:location|based\s+in|current\s+location|city)\s*[:\-]\s*([A-Za-z .,\-]{3,60})", re.I)
_NOTICE_RE = re.compile(
    r"(notice\s*period|availability|available\s+(?:to\s+join|from))\s*[:\-]?\s*([^\n|]{2,60})", re.I)

_INDIAN_CITIES = [
    "bengaluru", "bangalore", "mumbai", "delhi", "new delhi", "gurugram", "gurgaon", "noida",
    "hyderabad", "chennai", "pune", "kolkata", "ahmedabad", "jaipur", "kochi", "chandigarh",
    "indore", "coimbatore", "thiruvananthapuram", "remote",
]

_DEGREE_RE = re.compile(
    r"\b(ph\.?d|doctorate|m\.?tech|m\.?e\.?|m\.?s\.?c?|mba|pgdm|m\.?com|m\.?a\b|"
    r"b\.?tech|b\.?e\.?|b\.?sc|b\.?com|b\.?a\b|bba|bca|mca|diploma|bachelor[s']?|master[s']?|"
    r"post\s*graduate|under\s*graduate)\b", re.I)

_CERT_RE = re.compile(
    r"\b(certified|certification|certificate|credential|nanodegree|licensed)\b", re.I)
# "Notice Period: 30 days" and similar metadata often trail the last section.
_META_LINE_RE = re.compile(
    r"^\s*(notice\s*period|availability|available\s+from|current\s+ctc|expected\s+ctc|salary|"
    r"location|address|phone|mobile|email|languages?|references?|date\s+of\s+birth|linkedin|github)"
    r"\s*[:\-]", re.I)

_TITLE_DOMAIN = (
    r"(?:software|data|product|program|project|marketing|growth|performance|digital|business|sales|"
    r"machine\s+learning|ml|devops|cloud|qa|full[\s-]?stack|front[\s-]?end|back[\s-]?end|ui/ux|hr|"
    r"finance|content|brand|operations|research|security|platform|mobile|systems)"
)
_TITLE_RE = re.compile(
    r"\b((?:senior|sr\.?|junior|jr\.?|lead|principal|staff|chief|head\s+of|associate|assistant|deputy|vice)?\s*"
    + _TITLE_DOMAIN + r"?(?:\s+" + _TITLE_DOMAIN + r"){0,2}\s*"
    r"(?:engineer|developer|manager|analyst|scientist|architect|consultant|designer|specialist|lead|"
    r"director|executive|associate|officer|strategist|marketer|intern))\b", re.I)

# Words that look like a company when a line is split on delimiters, but are not.
_NOT_A_COMPANY = re.compile(
    r"^(present|current|now|today|ongoing|full[\s-]?time|part[\s-]?time|contract|intern(ship)?|"
    r"remote|onsite|hybrid|india|freelance|self[\s-]?employed)$", re.I)

# Note: "/" is deliberately NOT a delimiter — it would destroy A/B, CI/CD, UI/UX.
_SKILL_SPLIT_RE = re.compile(r"[,;|•●\t]|\s{3,}|\s-\s")

_ACTION_VERBS = re.compile(
    r"\b(led|managed|built|developed|designed|launched|owned|drove|delivered|scaled|improved|"
    r"increased|reduced|optimized|optimised|implemented|created|executed|ran|handled|coordinated|"
    r"mentored|analyzed|analysed|automated|migrated|architected|collaborated|partnered|grew|"
    r"achieved|generated|established|spearheaded|oversaw|supported|maintained)\b", re.I)


def split_sections(text: str) -> Dict[str, str]:
    """Split a resume into named sections using common headings.

    Layout-agnostic by design: a heading is any short line matching a known
    section name, so single-column and heading-per-block layouts both work.
    """
    lines = text.split("\n")
    sections: Dict[str, List[str]] = {}
    current = "header"
    for raw_line in lines:
        line = raw_line.strip()
        matched: Optional[str] = None
        if line and len(line) <= 60:
            probe = strip_bullets(line).strip(" :-–—")
            for name, pattern in SECTION_PATTERNS.items():
                if pattern.match(probe):
                    matched = name
                    break
        if matched:
            current = matched
            sections.setdefault(current, [])
            continue
        sections.setdefault(current, []).append(raw_line)
    return {k: normalize_text("\n".join(v)) for k, v in sections.items() if normalize_text("\n".join(v))}


def _extract_skills(sections: Dict[str, str], full_text: str) -> List[str]:
    blob = sections.get("skills", "")
    skills: List[str] = []
    seen: set[str] = set()

    for line in blob.split("\n"):
        line = strip_bullets(line).strip()
        if not line:
            continue
        # "Languages: Python, Go" -> keep only the value side
        if ":" in line and len(line.split(":", 1)[0]) < 40:
            line = line.split(":", 1)[1]
        for token in _SKILL_SPLIT_RE.split(line):
            token = token.strip(" .-–—")
            if 1 < len(token) <= 45 and not token.lower().startswith(("and ", "etc")):
                key = token.lower()
                if key not in seen:
                    seen.add(key)
                    skills.append(token)
    return skills[:80]


def _extract_education(sections: Dict[str, str], full_text: str) -> List[str]:
    blob = sections.get("education") or ""
    source = blob if blob else full_text
    out: List[str] = []
    seen: set[str] = set()
    for line in source.split("\n"):
        line = strip_bullets(line).strip()
        if not line or len(line) > 200:
            continue
        if _DEGREE_RE.search(line):
            key = line.lower()
            if key not in seen:
                seen.add(key)
                out.append(line)
    return out[:12]


def _extract_certifications(sections: Dict[str, str], full_text: str) -> List[str]:
    blob = sections.get("certifications") or ""
    out: List[str] = []
    seen: set[str] = set()
    for line in (blob or "").split("\n"):
        line = strip_bullets(line).strip()
        if _META_LINE_RE.match(line):
            continue
        if 3 < len(line) <= 160:
            key = line.lower()
            if key not in seen:
                seen.add(key)
                out.append(line)
    if not out:
        # No certifications section: scan the whole resume, but only accept
        # credential-shaped lines. A sentence that merely uses the word
        # "certification" is a responsibility, not a credential.
        for line in full_text.split("\n"):
            line = strip_bullets(line).strip()
            if not _CERT_RE.search(line) or not 3 < len(line) <= 110:
                continue
            if _ACTION_VERBS.search(line) or line.rstrip().endswith("."):
                continue
            key = line.lower()
            if key not in seen:
                seen.add(key)
                out.append(line)
    return out[:15]


def _extract_titles_and_companies(sections: Dict[str, str], full_text: str) -> tuple[List[str], List[str]]:
    blob = sections.get("experience") or full_text
    titles: List[str] = []
    companies: List[str] = []
    seen_t: set[str] = set()
    seen_c: set[str] = set()

    for raw_line in blob.split("\n"):
        line = strip_bullets(raw_line).strip()
        if not line or len(line) > 160:
            continue
        match = _TITLE_RE.search(line)
        if not match:
            continue
        title = re.sub(r"\s+", " ", match.group(1)).strip().title()
        if title.lower() not in seen_t:
            seen_t.add(title.lower())
            titles.append(title)

        # Companies usually sit next to the title, separated by a delimiter.
        for part in re.split(r"\s+(?:at|@|\||,|-|–|—)\s+", line):
            part = part.strip()
            if not part or _TITLE_RE.fullmatch(part) or _TITLE_RE.search(part):
                continue
            part = re.sub(r"\(.*?\)", "", part).strip(" .,-|")
            part = re.sub(r"\b(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?\s*\d{4}.*$", "", part, flags=re.I).strip(" .,-|")
            if _NOT_A_COMPANY.match(part):
                continue
            if 2 < len(part) <= 60 and not re.fullmatch(r"[\d\s./-]+", part) and part.lower() not in seen_c:
                seen_c.add(part.lower())
                companies.append(part)
    return titles[:15], companies[:15]


def _extract_responsibilities(sections: Dict[str, str], full_text: str) -> List[str]:
    blob = "\n".join(v for k, v in sections.items() if k in {"experience", "projects", "achievements", "summary"})
    source = blob if blob else full_text
    out: List[str] = []
    for unit in split_evidence_units(source):
        if _ACTION_VERBS.search(unit):
            out.append(unit)
    return out[:60]


def _extract_contact(text: str) -> tuple[str, str, str, str]:
    email_match = _EMAIL_RE.search(text)
    phone_match = _PHONE_RE.search(text)
    email = email_match.group(0) if email_match else ""
    phone = phone_match.group(0).strip() if phone_match else ""

    location = ""
    loc_match = _LOCATION_RE.search(text)
    if loc_match:
        location = loc_match.group(1).strip(" .,-")
    else:
        head = "\n".join(text.split("\n")[:12]).lower()
        for city in _INDIAN_CITIES:
            if re.search(rf"(?<![a-z]){re.escape(city)}(?![a-z])", head):
                location = city.title()
                break

    notice = ""
    notice_match = _NOTICE_RE.search(text)
    if notice_match:
        notice = notice_match.group(2).strip(" .,-")
    return email, phone, location, notice


# ---------------------------------------------------------------------------
# Public entry point
# ---------------------------------------------------------------------------

def parse_resume(file_name: str, data: bytes, fallback_name: str = "Candidate") -> ParsedResume:
    """Parse one resume file. Always returns a ParsedResume — never raises."""
    result = ParsedResume(file_name=file_name)
    result.content_sha = content_hash(data or b"")

    try:
        raw_text, pages, method = extract_text(file_name, data)
    except ResumeParseError as exc:
        result.status = exc.status
        result.error = str(exc)
        result.candidate_name = fallback_name
        return result
    except Exception as exc:  # noqa: BLE001 - defensive: one file must not kill a batch
        logger.exception("Unexpected parse failure for %s", file_name)
        result.status = "failed"
        result.error = f"Unexpected error while reading the file: {str(exc)[:160]}"
        result.candidate_name = fallback_name
        return result

    # Re-join lines that the source layout split mid-sentence. PDF extraction and
    # OCR both wrap long bullets across several lines, which would otherwise cut
    # evidence quotes off halfway through.
    text = dewrap_lines(normalize_text(raw_text))
    if len(text) < MIN_TOTAL_CHARS:
        result.status = "failed"
        result.error = "Empty resume — no usable text could be extracted."
        result.candidate_name = fallback_name
        result.text = text
        result.extraction_method = method
        return result

    result.text = text
    result.page_count = pages
    result.char_count = len(text)
    result.extraction_method = method
    result.text_sha = text_fingerprint(text)

    name, extracted = extract_candidate_name(text, fallback_name)
    result.candidate_name = name
    result.name_extracted = extracted

    result.email, result.phone, result.location, result.notice_period = _extract_contact(text)
    result.sections = split_sections(text)
    result.skills = _extract_skills(result.sections, text)
    result.education = _extract_education(result.sections, text)
    result.certifications = _extract_certifications(result.sections, text)
    result.job_titles, result.companies = _extract_titles_and_companies(result.sections, text)
    result.responsibilities = _extract_responsibilities(result.sections, text)
    result.evidence_units = split_evidence_units(text)

    stated = extract_total_years(text)
    derived = estimate_years_from_dates(text)
    if stated is not None:
        result.total_years = stated
        result.years_source = "stated in resume"
    elif derived is not None:
        result.total_years = derived
        result.years_source = "derived from employment dates in resume"
    else:
        result.years_source = NOT_MENTIONED

    return result


def find_duplicates(parsed: List[ParsedResume]) -> Dict[str, str]:
    """Map file_name -> the earlier file it duplicates.

    Matches on identical bytes first, then on a normalised-text fingerprint so
    the same resume re-exported to another format is still caught.
    """
    by_content: Dict[str, str] = {}
    by_text: Dict[str, str] = {}
    duplicates: Dict[str, str] = {}

    for item in parsed:
        if not item.ok:
            continue
        original = by_content.get(item.content_sha) or by_text.get(item.text_sha)
        if original:
            duplicates[item.file_name] = original
            continue
        by_content[item.content_sha] = item.file_name
        by_text[item.text_sha] = item.file_name
    return duplicates
