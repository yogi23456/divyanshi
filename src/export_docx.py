"""DOCX export: a professional candidate evaluation report.

Structure:
  1. Cover block — role, batch statistics, how scoring was configured
  2. Ranking table — every processed candidate
  3. Per-candidate sections — score breakdown with reasoning, strengths,
     concerns, mandatory requirements, missing information, resume evidence and
     a recruiter notes block
  4. Files that could not be processed, with the reason for each
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Dict, Optional, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .config import NOT_MENTIONED, SCORE_CATEGORIES
from .jd_parser import ParsedJD
from .scoring import CandidateEvaluation
from .utils import truncate

BRAND = RGBColor(0x1F, 0x38, 0x64)
MUTED = RGBColor(0x59, 0x59, 0x59)

RECOMMENDATION_COLORS = {
    "Strong Match": "C6EFCE",
    "Match": "DDEBF7",
    "Partial Match": "FFF2CC",
    "Weak Match": "FCE4D6",
    "Insufficient Information": "E7E6E6",
}


def _shade(cell, hex_color: str) -> None:
    """Apply a background colour to a table cell."""
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shading)


def _para(document, text: str = "", *, size: int = 10, bold: bool = False,
          italic: bool = False, color: Optional[RGBColor] = None, space_after: int = 4,
          align=None):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    return paragraph


def _bullets(document, items: Sequence[str], empty_text: str = "None recorded.") -> None:
    if not items:
        _para(document, empty_text, size=9.5, italic=True, color=MUTED)
        return
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(str(item))
        run.font.size = Pt(9.5)


def _heading(document, text: str, level: int = 1) -> None:
    heading = document.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = BRAND


def _style_table_header(row) -> None:
    for cell in row.cells:
        _shade(cell, "1F3864")
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def _set_cell(cell, text: str, size: float = 9, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.font.size = Pt(size)
    run.bold = bold


# ---------------------------------------------------------------------------

def _write_cover(document, summary: Dict[str, object], jd: Optional[ParsedJD],
                 weights: Dict[str, float], analysis_method: str) -> None:
    title = document.add_heading("Candidate Evaluation Report", level=0)
    for run in title.runs:
        run.font.color.rgb = BRAND

    if jd is not None and jd.title:
        _para(document, jd.title, size=13, bold=True)
    _para(document,
          f"Generated {datetime.now():%d %B %Y, %H:%M}  ·  Analysis method: {analysis_method}",
          size=9, italic=True, color=MUTED, space_after=10)

    if jd is not None:
        _heading(document, "Job Description", level=1)
        table = document.add_table(rows=0, cols=2)
        table.style = "Light List Accent 1"
        for label, value in [
            ("Source file", jd.file_name or "-"),
            ("Minimum experience", f"{jd.min_years:g} years" if jd.min_years is not None else "Not stated"),
            ("Mandatory requirements", str(len(jd.mandatory()))),
            ("Preferred requirements", str(len(jd.preferred()))),
            ("Location requirement", jd.location_requirements[0] if jd.location_requirements
             else "Not stated in the JD"),
        ]:
            cells = table.add_row().cells
            _set_cell(cells[0], label, bold=True)
            _set_cell(cells[1], truncate(value, 220))

    _heading(document, "Batch Summary", level=1)
    counts: Dict[str, int] = summary.get("counts", {})  # type: ignore[assignment]
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header = table.rows[0]
    for index, title_text in enumerate(["Metric", "Value", "Recommendation", "Candidates"]):
        _set_cell(header.cells[index], title_text, bold=True)
    _style_table_header(header)

    metrics = [
        ("Total resumes uploaded", summary.get("total_uploaded", 0)),
        ("Successfully processed", summary.get("processed", 0)),
        ("Failed / OCR required", summary.get("failed", 0)),
        ("Duplicates detected", summary.get("duplicates", 0)),
        ("Average score", summary.get("average_score", 0)),
        ("Highest score", summary.get("highest_score", 0)),
        ("Lowest score", summary.get("lowest_score", 0)),
    ]
    recommendation_rows = list(counts.items())
    for index in range(max(len(metrics), len(recommendation_rows))):
        cells = table.add_row().cells
        if index < len(metrics):
            _set_cell(cells[0], metrics[index][0])
            _set_cell(cells[1], metrics[index][1])
        if index < len(recommendation_rows):
            name, count = recommendation_rows[index]
            _set_cell(cells[2], name)
            _set_cell(cells[3], count)
            _shade(cells[2], RECOMMENDATION_COLORS.get(name, "E7E6E6"))

    _heading(document, "Scoring Weights", level=1)
    _para(document,
          "These weights were configured by the recruiter for this run and total "
          f"{round(sum(weights.values()), 2):g} points.", size=9.5, color=MUTED)
    weights_table = document.add_table(rows=1, cols=2)
    weights_table.style = "Light List Accent 1"
    _set_cell(weights_table.rows[0].cells[0], "Category", bold=True)
    _set_cell(weights_table.rows[0].cells[1], "Maximum points", bold=True)
    for key, label in SCORE_CATEGORIES.items():
        cells = weights_table.add_row().cells
        _set_cell(cells[0], label)
        _set_cell(cells[1], f"{weights.get(key, 0):g}")


def _write_ranking_table(document, evaluations: Sequence[CandidateEvaluation]) -> None:
    _heading(document, "Overall Ranking", level=1)
    processed = [e for e in evaluations if e.ok]
    if not processed:
        _para(document, "No resumes were successfully processed in this batch.",
              size=10, italic=True, color=MUTED)
        return

    table = document.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    header = table.rows[0]
    for index, title in enumerate(
            ["#", "Candidate", "Score", "Recommendation", "Mandatory Met", "Experience"]):
        _set_cell(header.cells[index], title, bold=True)
    _style_table_header(header)

    for evaluation in processed:
        cells = table.add_row().cells
        _set_cell(cells[0], evaluation.rank)
        _set_cell(cells[1], evaluation.candidate_name)
        _set_cell(cells[2], f"{evaluation.overall_score:g}/100", bold=True)
        _set_cell(cells[3], evaluation.recommendation)
        _shade(cells[3], RECOMMENDATION_COLORS.get(evaluation.recommendation, "E7E6E6"))
        total_mandatory = evaluation.mandatory_met_count() + evaluation.mandatory_missed_count()
        _set_cell(cells[4], f"{evaluation.mandatory_met_count()}/{total_mandatory}")
        _set_cell(cells[5], f"{evaluation.total_years:g} yrs" if evaluation.total_years is not None
                  else NOT_MENTIONED)

    for column, width in zip(table.columns, [0.4, 1.9, 0.8, 1.4, 1.0, 1.0]):
        for cell in column.cells:
            cell.width = Inches(width)


def _write_candidate(document, evaluation: CandidateEvaluation) -> None:
    _heading(document, f"{evaluation.rank}. {evaluation.candidate_name}", level=1)
    _para(document, f"Resume: {evaluation.file_name}", size=9, italic=True, color=MUTED, space_after=2)
    _para(document,
          f"Overall score: {evaluation.overall_score:g}/100   ·   Recommendation: {evaluation.recommendation}",
          size=11, bold=True, space_after=2)
    _para(document, f"Why this recommendation: {evaluation.recommendation_reason}",
          size=9.5, color=MUTED, space_after=8)

    _heading(document, "Score Breakdown", level=2)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    header = table.rows[0]
    for index, title in enumerate(["Category", "Score", "Reasoning"]):
        _set_cell(header.cells[index], title, bold=True)
    _style_table_header(header)
    for category in evaluation.category_scores:
        cells = table.add_row().cells
        _set_cell(cells[0], category.label, bold=True)
        _set_cell(cells[1], category.display())
        _set_cell(cells[2], category.reasoning, size=8.5)
    total_row = table.add_row().cells
    _set_cell(total_row[0], "TOTAL", bold=True)
    _set_cell(total_row[1], f"{evaluation.overall_score:g}/100", bold=True)
    _set_cell(total_row[2], "", bold=True)
    for cell in total_row:
        _shade(cell, "F2F2F2")
    for column, width in zip(table.columns, [1.7, 0.9, 4.0]):
        for cell in column.cells:
            cell.width = Inches(width)

    _heading(document, "Relevant Experience", level=2)
    _para(document, evaluation.relevant_experience, size=9.5)

    _heading(document, "Key Strengths", level=2)
    _bullets(document, evaluation.key_strengths, "No standout strengths were evidenced.")

    _heading(document, "Key Concerns", level=2)
    _bullets(document, evaluation.key_concerns, "No material concerns were identified.")

    _heading(document, "Mandatory Requirements", level=2)
    _para(document, "Met:", size=9.5, bold=True, space_after=2)
    _bullets(document, evaluation.mandatory_requirements_met, "None evidenced.")
    _para(document, "Not evidenced:", size=9.5, bold=True, space_after=2)
    _bullets(document, evaluation.mandatory_requirements_missed, "None — all mandatory requirements are met.")

    _heading(document, "Missing Information", level=2)
    _para(document,
          "Recorded exactly as found: nothing below is inferred about the candidate.",
          size=8.5, italic=True, color=MUTED, space_after=3)
    _bullets(document, evaluation.missing_information, "Nothing material was missing from this resume.")

    _heading(document, "Resume Evidence", level=2)
    _para(document, "Quotes taken from the resume, tagged with the requirement they support.",
          size=8.5, italic=True, color=MUTED, space_after=3)
    _bullets(document, [f'{item}' for item in evaluation.evidence],
             "No direct evidence quotes were extracted.")

    _heading(document, "Recruiter Notes", level=2)
    notes_table = document.add_table(rows=0, cols=2)
    notes_table.style = "Light List Accent 1"
    for label, value in [
        ("Shortlisted", "Yes" if evaluation.shortlisted else "No"),
        ("Interview status", evaluation.interview_status),
        ("Final decision", evaluation.recruiter_decision or "Not recorded"),
        ("Notes", evaluation.recruiter_notes or "________________________________________"),
    ]:
        cells = notes_table.add_row().cells
        _set_cell(cells[0], label, bold=True)
        _set_cell(cells[1], value)
        cells[0].width = Inches(1.4)
        cells[1].width = Inches(5.2)


def _write_unprocessed(document, evaluations: Sequence[CandidateEvaluation]) -> None:
    unprocessed = [e for e in evaluations if not e.ok]
    if not unprocessed:
        return
    document.add_page_break()
    _heading(document, "Files Not Evaluated", level=1)
    _para(document,
          "These uploads were not scored. Each one is listed with the reason so nothing "
          "disappears silently from the batch.", size=9.5, color=MUTED)

    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    header = table.rows[0]
    for index, title in enumerate(["File", "Status", "Reason"]):
        _set_cell(header.cells[index], title, bold=True)
    _style_table_header(header)
    for evaluation in unprocessed:
        cells = table.add_row().cells
        _set_cell(cells[0], evaluation.file_name)
        _set_cell(cells[1], evaluation.status)
        _set_cell(cells[2], evaluation.recommendation_reason or evaluation.error or "-", size=8.5)
    for column, width in zip(table.columns, [2.2, 1.2, 3.2]):
        for cell in column.cells:
            cell.width = Inches(width)


def export_to_docx(
    evaluations: Sequence[CandidateEvaluation],
    summary: Dict[str, object],
    weights: Dict[str, float],
    output_path: Path | str,
    jd: Optional[ParsedJD] = None,
    analysis_method: str = "evidence-engine",
    max_candidates: Optional[int] = None,
) -> Path:
    """Write the candidate evaluation report and return its path."""
    document = Document()
    for section in document.sections:
        section.left_margin = section.right_margin = Inches(0.8)
        section.top_margin = section.bottom_margin = Inches(0.7)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)

    _write_cover(document, summary, jd, weights, analysis_method)
    document.add_page_break()
    _write_ranking_table(document, evaluations)

    processed = [e for e in evaluations if e.ok]
    if max_candidates is not None:
        processed = processed[:max_candidates]

    if processed:
        document.add_page_break()
        _heading(document, "Detailed Candidate Evaluations", level=1)
        for index, evaluation in enumerate(processed):
            if index:
                document.add_page_break()
            _write_candidate(document, evaluation)

    _write_unprocessed(document, evaluations)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    return output_path
