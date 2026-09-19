"""Excel export: a three-sheet workbook a recruiter can actually work in.

  Candidate Ranking    one row per candidate, colour-coded, filterable
  Detailed Evaluation  one row per candidate per scoring category, with reasoning
                       and the resume evidence behind it
  Summary              batch statistics and the top candidates

Formatting is deliberate: frozen headers, autofilter, wrapped text, sensible
column widths and colour scales, so a 500-row sheet stays scannable.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Sequence

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .config import NOT_MENTIONED, SCORE_CATEGORIES
from .jd_parser import ParsedJD
from .scoring import CandidateEvaluation
from .utils import truncate

# --- Palette ---------------------------------------------------------------
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F3864")
SUBTLE_FONT = Font(italic=True, size=9, color="595959")
SECTION_FONT = Font(bold=True, size=11, color="1F3864")
BORDER = Border(*[Side(style="thin", color="D9D9D9")] * 4)

RECOMMENDATION_FILLS = {
    "Strong Match": PatternFill("solid", fgColor="C6EFCE"),
    "Match": PatternFill("solid", fgColor="DDEBF7"),
    "Partial Match": PatternFill("solid", fgColor="FFF2CC"),
    "Weak Match": PatternFill("solid", fgColor="FCE4D6"),
    "Insufficient Information": PatternFill("solid", fgColor="E7E6E6"),
}

RANKING_COLUMNS: List[tuple[str, int]] = [
    ("Rank", 7),
    ("Candidate Name", 24),
    ("Resume", 30),
    ("Overall Score", 13),
    ("Recommendation", 22),
    ("Relevant Experience", 40),
    ("Required Skills Match", 18),
    ("JD Alignment", 14),
    ("Industry Match", 15),
    ("Education", 12),
    ("Preferred Skills", 15),
    ("Mandatory Requirements Missed", 46),
    ("Key Strengths", 60),
    ("Key Concerns", 60),
    ("Years of Experience", 13),
    ("Location", 18),
    ("Notice Period", 15),
    ("Missing Information", 46),
    ("Shortlisted", 11),
    ("Interview Status", 18),
    ("Recruiter Notes", 40),
    ("Recruiter Final Decision", 24),
    ("Processing Status", 16),
]

DETAIL_COLUMNS: List[tuple[str, int]] = [
    ("Candidate", 24),
    ("Resume", 28),
    ("Category", 26),
    ("Score", 9),
    ("Max Score", 10),
    ("Score %", 9),
    ("Reasoning", 78),
    ("Resume Evidence", 78),
]


def _style_header(sheet: Worksheet, columns: Sequence[tuple[str, int]], row: int = 1) -> None:
    for index, (title, width) in enumerate(columns, start=1):
        cell = sheet.cell(row=row, column=index, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        cell.border = BORDER
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[row].height = 30


def _finish_sheet(sheet: Worksheet, columns: Sequence[tuple[str, int]],
                  header_row: int, last_row: int) -> None:
    """Freeze the header, enable filters and wrap the long text columns."""
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
    if last_row > header_row:
        sheet.auto_filter.ref = (
            f"A{header_row}:{get_column_letter(len(columns))}{last_row}")
    for row in sheet.iter_rows(min_row=header_row + 1, max_row=last_row,
                               max_col=len(columns)):
        for cell in row:
            cell.border = BORDER
            if cell.alignment.horizontal is None:
                cell.alignment = Alignment(vertical="top", wrap_text=True)


def _join(items: Sequence[str], bullet: str = "• ") -> str:
    return "\n".join(f"{bullet}{item}" for item in items) if items else ""


# ---------------------------------------------------------------------------
# Sheet 1 — Candidate Ranking
# ---------------------------------------------------------------------------

def _write_ranking(sheet: Worksheet, evaluations: Sequence[CandidateEvaluation]) -> None:
    _style_header(sheet, RANKING_COLUMNS)
    row = 2
    for evaluation in evaluations:
        values = [
            evaluation.rank or "",
            evaluation.candidate_name,
            evaluation.file_name,
            evaluation.overall_score if evaluation.ok else "",
            evaluation.recommendation,
            evaluation.relevant_experience,
            evaluation.display_score("required_skills"),
            evaluation.display_score("jd_alignment"),
            evaluation.display_score("industry_relevance"),
            evaluation.display_score("education"),
            evaluation.display_score("preferred_skills"),
            _join(evaluation.mandatory_requirements_missed) or "None",
            _join(evaluation.key_strengths),
            _join(evaluation.key_concerns),
            evaluation.total_years if evaluation.total_years is not None else NOT_MENTIONED,
            evaluation.location,
            evaluation.notice_period,
            _join(evaluation.missing_information),
            "Yes" if evaluation.shortlisted else "No",
            evaluation.interview_status,
            evaluation.recruiter_notes,
            evaluation.recruiter_decision,
            evaluation.status if evaluation.ok else f"{evaluation.status}: {truncate(evaluation.error, 90)}",
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

        sheet.cell(row=row, column=1).alignment = Alignment(horizontal="center", vertical="top")
        score_cell = sheet.cell(row=row, column=4)
        score_cell.alignment = Alignment(horizontal="center", vertical="top")
        score_cell.font = Font(bold=True)
        score_cell.number_format = "0.0"

        recommendation_cell = sheet.cell(row=row, column=5)
        recommendation_cell.fill = RECOMMENDATION_FILLS.get(
            evaluation.recommendation, RECOMMENDATION_FILLS["Insufficient Information"])
        recommendation_cell.font = Font(bold=True)
        recommendation_cell.alignment = Alignment(horizontal="center", vertical="top")

        if not evaluation.ok:
            for column in range(1, len(RANKING_COLUMNS) + 1):
                sheet.cell(row=row, column=column).font = Font(color="808080", italic=True)
        row += 1

    last_row = row - 1
    _finish_sheet(sheet, RANKING_COLUMNS, 1, last_row)

    if last_row >= 2:
        # Green-amber-red scale on the score column so the eye lands on the top.
        sheet.conditional_formatting.add(
            f"D2:D{last_row}",
            ColorScaleRule(start_type="num", start_value=0, start_color="F8696B",
                           mid_type="num", mid_value=55, mid_color="FFEB84",
                           end_type="num", end_value=100, end_color="63BE7B"))
        # Flag any candidate with an unmet mandatory requirement.
        sheet.conditional_formatting.add(
            f"L2:L{last_row}",
            CellIsRule(operator="notEqual", formula=['"None"'],
                       fill=PatternFill("solid", fgColor="FCE4D6")))


# ---------------------------------------------------------------------------
# Sheet 2 — Detailed Evaluation
# ---------------------------------------------------------------------------

def _write_detail(sheet: Worksheet, evaluations: Sequence[CandidateEvaluation]) -> None:
    _style_header(sheet, DETAIL_COLUMNS)
    row = 2
    for evaluation in evaluations:
        if not evaluation.ok:
            for column, value in enumerate(
                [evaluation.candidate_name, evaluation.file_name, "Not evaluated", "", "", "",
                 evaluation.recommendation_reason or evaluation.error, ""], start=1):
                cell = sheet.cell(row=row, column=column, value=value)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.font = Font(color="808080", italic=True)
            row += 1
            continue

        for category in evaluation.category_scores:
            values = [
                evaluation.candidate_name,
                evaluation.file_name,
                category.label,
                category.score,
                category.max_score,
                category.percentage / 100.0,
                category.reasoning,
                _join([f'"{item}"' for item in category.evidence]) or "No direct quote available for this category.",
            ]
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row=row, column=column, value=value)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            sheet.cell(row=row, column=4).number_format = "0.0"
            sheet.cell(row=row, column=5).number_format = "0.0"
            sheet.cell(row=row, column=6).number_format = "0%"
            row += 1

        # A per-candidate mandatory-requirements roll-up closes each block.
        for column, value in enumerate([
            evaluation.candidate_name, evaluation.file_name, "Mandatory Requirements",
            evaluation.mandatory_met_count(),
            evaluation.mandatory_met_count() + evaluation.mandatory_missed_count(),
            evaluation.mandatory_coverage(),
            "Met: " + (_join(evaluation.mandatory_requirements_met) or "None")
            + "\nMissed: " + (_join(evaluation.mandatory_requirements_missed) or "None"),
            _join([f'"{item}"' for item in evaluation.evidence[:6]]),
        ], start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.font = Font(bold=True)
        sheet.cell(row=row, column=6).number_format = "0%"
        row += 1

    last_row = row - 1
    _finish_sheet(sheet, DETAIL_COLUMNS, 1, last_row)
    if last_row >= 2:
        sheet.conditional_formatting.add(
            f"F2:F{last_row}",
            ColorScaleRule(start_type="num", start_value=0, start_color="F8696B",
                           mid_type="num", mid_value=0.55, mid_color="FFEB84",
                           end_type="num", end_value=1, end_color="63BE7B"))


# ---------------------------------------------------------------------------
# Sheet 3 — Summary
# ---------------------------------------------------------------------------

def _write_summary(sheet: Worksheet, summary: Dict[str, object], jd: Optional[ParsedJD],
                   weights: Dict[str, float], analysis_method: str) -> None:
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 20
    sheet.column_dimensions["C"].width = 26
    sheet.column_dimensions["D"].width = 26
    sheet.column_dimensions["E"].width = 34

    sheet["A1"] = "Candidate Evaluation Summary"
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = f"Generated {datetime.now():%d %b %Y, %H:%M}  ·  Analysis: {analysis_method}"
    sheet["A2"].font = SUBTLE_FONT

    row = 4
    if jd is not None:
        sheet[f"A{row}"] = "Job Description"
        sheet[f"A{row}"].font = SECTION_FONT
        row += 1
        for label, value in [
            ("Role", jd.title or "Not stated"),
            ("Source file", jd.file_name or "-"),
            ("Minimum experience required", f"{jd.min_years:g} years" if jd.min_years is not None else "Not stated"),
            ("Mandatory requirements", len(jd.mandatory())),
            ("Preferred requirements", len(jd.preferred())),
        ]:
            sheet[f"A{row}"] = label
            sheet[f"A{row}"].font = Font(bold=True)
            sheet[f"B{row}"] = value
            row += 1
        row += 1

    sheet[f"A{row}"] = "Processing"
    sheet[f"A{row}"].font = SECTION_FONT
    row += 1
    for label, value in [
        ("Total resumes uploaded", summary.get("total_uploaded", 0)),
        ("Successfully processed", summary.get("processed", 0)),
        ("Failed / OCR required", summary.get("failed", 0)),
        ("Duplicates detected", summary.get("duplicates", 0)),
    ]:
        sheet[f"A{row}"] = label
        sheet[f"A{row}"].font = Font(bold=True)
        sheet[f"B{row}"] = value
        row += 1
    row += 1

    sheet[f"A{row}"] = "Recommendation breakdown"
    sheet[f"A{row}"].font = SECTION_FONT
    row += 1
    counts: Dict[str, int] = summary.get("counts", {})  # type: ignore[assignment]
    for name, count in counts.items():
        sheet[f"A{row}"] = name
        sheet[f"A{row}"].font = Font(bold=True)
        sheet[f"A{row}"].fill = RECOMMENDATION_FILLS.get(name, RECOMMENDATION_FILLS["Insufficient Information"])
        sheet[f"B{row}"] = count
        row += 1
    row += 1

    sheet[f"A{row}"] = "Scores"
    sheet[f"A{row}"].font = SECTION_FONT
    row += 1
    for label, value in [
        ("Average candidate score", summary.get("average_score", 0)),
        ("Highest score", summary.get("highest_score", 0)),
        ("Lowest score", summary.get("lowest_score", 0)),
        ("Shortlisted by recruiter", summary.get("shortlisted", 0)),
    ]:
        sheet[f"A{row}"] = label
        sheet[f"A{row}"].font = Font(bold=True)
        sheet[f"B{row}"] = value
        sheet[f"B{row}"].number_format = "0.0"
        row += 1
    row += 1

    sheet[f"A{row}"] = "Scoring weights used"
    sheet[f"A{row}"].font = SECTION_FONT
    row += 1
    for key, label in SCORE_CATEGORIES.items():
        sheet[f"A{row}"] = label
        sheet[f"A{row}"].font = Font(bold=True)
        sheet[f"B{row}"] = weights.get(key, 0)
        row += 1
    sheet[f"A{row}"] = "Total"
    sheet[f"A{row}"].font = Font(bold=True)
    sheet[f"B{row}"] = round(sum(weights.values()), 2)
    sheet[f"B{row}"].font = Font(bold=True)
    row += 2

    sheet[f"A{row}"] = "Top candidates"
    sheet[f"A{row}"].font = SECTION_FONT
    row += 1
    for index, (title, width) in enumerate(
            [("Rank", 7), ("Candidate", 26), ("Score", 10), ("Recommendation", 24), ("Resume", 34)], start=1):
        cell = sheet.cell(row=row, column=index, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    row += 1
    for candidate in summary.get("top_candidates", []):  # type: ignore[union-attr]
        sheet.cell(row=row, column=1, value=candidate["rank"]).alignment = Alignment(horizontal="center")
        sheet.cell(row=row, column=2, value=candidate["name"])
        score_cell = sheet.cell(row=row, column=3, value=candidate["score"])
        score_cell.number_format = "0.0"
        score_cell.alignment = Alignment(horizontal="center")
        recommendation_cell = sheet.cell(row=row, column=4, value=candidate["recommendation"])
        recommendation_cell.fill = RECOMMENDATION_FILLS.get(
            candidate["recommendation"], RECOMMENDATION_FILLS["Insufficient Information"])
        sheet.cell(row=row, column=5, value=candidate["file"])
        row += 1

    row += 1
    sheet[f"A{row}"] = ("Scores are evidence-based: see the 'Detailed Evaluation' sheet for the "
                        "reasoning and the resume quote behind every category score.")
    sheet[f"A{row}"].font = SUBTLE_FONT


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

def export_to_excel(
    evaluations: Sequence[CandidateEvaluation],
    summary: Dict[str, object],
    weights: Dict[str, float],
    output_path: Path | str,
    jd: Optional[ParsedJD] = None,
    analysis_method: str = "evidence-engine",
) -> Path:
    """Write the three-sheet recruiter workbook and return its path."""
    workbook = Workbook()

    ranking_sheet = workbook.active
    ranking_sheet.title = "Candidate Ranking"
    _write_ranking(ranking_sheet, evaluations)

    _write_detail(workbook.create_sheet("Detailed Evaluation"), evaluations)
    _write_summary(workbook.create_sheet("Summary"), summary, jd, weights, analysis_method)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(str(output_path))
    return output_path
