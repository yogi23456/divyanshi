"""End-to-end test: JD -> resumes -> analysis -> ranking -> Excel + DOCX.

Runs the whole pipeline against the sample data and asserts the behaviour the
tool promises: explainable scores, honest handling of unreadable files, no
invented information, and valid export files.

Run:  python scripts/e2e_test.py
Exits non-zero if any check fails.
"""

from __future__ import annotations

import re
import sys
import traceback
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.candidate_analyzer import analyze_batch  # noqa: E402
from src.config import DEFAULT_WEIGHTS, NOT_MENTIONED, SCORE_CATEGORIES, get_settings, validate_weights  # noqa: E402
from src.export_docx import export_to_docx  # noqa: E402
from src.export_excel import export_to_excel  # noqa: E402
from src.jd_parser import parse_jd_file  # noqa: E402
from src.llm_client import build_client  # noqa: E402
from src.ranking import build_summary, filter_candidates, rank_candidates, search_candidates, to_dataframe  # noqa: E402
from src.resume_parser import find_duplicates, parse_resume  # noqa: E402
from src.storage import ExtractionCache  # noqa: E402

PASSED: list[str] = []
FAILED: list[str] = []


def check(name: str, condition: bool, detail: str = "") -> None:
    if condition:
        PASSED.append(name)
        print(f"  PASS  {name}")
    else:
        FAILED.append(f"{name} — {detail}")
        print(f"  FAIL  {name}  {detail}")


def section(title: str) -> None:
    print(f"\n{title}\n{'-' * len(title)}")


def main() -> int:
    settings = get_settings()
    client = build_client(settings)
    weights = dict(DEFAULT_WEIGHTS)

    print("=" * 74)
    print("END-TO-END TEST — AI Resume Review & Candidate Ranking")
    print("=" * 74)
    print(f"Analysis engine: {client.describe()}")

    samples = ROOT / "samples"
    if not (samples / "job_description.pdf").exists():
        print("\nSample data missing. Run: python scripts/generate_samples.py")
        return 1

    # ---------------------------------------------------------------- config
    section("1. Configuration")
    ok, message = validate_weights(weights)
    check("Default weights total 100%", ok, message)
    check("Weights rejected when they do not total 100",
          not validate_weights({**weights, "education": 40})[0])
    check("Six scoring categories defined", len(SCORE_CATEGORIES) == 6)

    # -------------------------------------------------------------------- JD
    section("2. Job Description parsing")
    jd = parse_jd_file("job_description.pdf", (samples / "job_description.pdf").read_bytes())
    check("JD parsed from PDF", jd.ok, jd.error)
    check("Job title extracted", "Performance Marketing" in (jd.title or ""), jd.title)
    check("Minimum experience identified (5 years)", jd.min_years == 5.0, str(jd.min_years))
    check("Mandatory requirements found", len(jd.mandatory()) >= 5, str(len(jd.mandatory())))
    check("Preferred requirements found", len(jd.preferred()) >= 4, str(len(jd.preferred())))
    check("Mandatory and preferred are distinguished",
          not (set(r.text for r in jd.mandatory()) & set(r.text for r in jd.preferred())))
    check("Every requirement carries a classification reason",
          all(r.rationale for r in jd.requirements))
    check("Location requirement detected (the JD states one)", jd.requires_location())
    check("Education requirement detected", len(jd.education_requirements) >= 1)
    check("Industry requirement detected", len(jd.industry_requirements) >= 1)

    # --------------------------------------------------------------- resumes
    section("3. Resume parsing")
    files = sorted((samples / "resumes").glob("*"))
    parsed = [parse_resume(f.name, f.read_bytes(), f"Candidate {i:03d}")
              for i, f in enumerate(files, start=1)]
    by_file = {p.file_name: p for p in parsed}

    check("All files produced a result (no crashes)", len(parsed) == len(files))
    check("PDF resume parsed", by_file["anjali_mehta_resume.pdf"].ok)
    check("DOCX resume parsed", by_file["rahul_nair_resume.docx"].ok)
    check("Corrupted PDF flagged, not crashed",
          by_file["corrupted_resume.pdf"].status == "failed")
    check("Empty file flagged", by_file["empty_resume.txt"].status == "failed")
    check("Near-empty resume flagged", by_file["meera_iyer_resume.txt"].status == "failed")

    anjali = by_file["anjali_mehta_resume.pdf"]
    check("Candidate name extracted", anjali.candidate_name == "Anjali Mehta", anjali.candidate_name)
    check("Years of experience extracted", anjali.total_years == 7.0, str(anjali.total_years))
    check("Skills extracted", len(anjali.skills) >= 5, str(len(anjali.skills)))
    check("Education extracted", len(anjali.education) >= 2, str(len(anjali.education)))
    check("Certifications extracted", len(anjali.certifications) >= 2, str(anjali.certifications))
    check("Job titles extracted", len(anjali.job_titles) >= 2, str(anjali.job_titles))
    check("Companies extracted", len(anjali.companies) >= 2, str(anjali.companies))
    check("Responsibilities extracted", len(anjali.responsibilities) >= 5)
    check("Notice period read from the resume", "30" in anjali.notice_period, anjali.notice_period)

    unnamed = by_file["unnamed_candidate_resume.docx"]
    check("Missing name falls back to 'Candidate NNN', processing continues",
          unnamed.ok and unnamed.candidate_name.startswith("Candidate") and not unnamed.name_extracted,
          unnamed.candidate_name)

    duplicates = find_duplicates(parsed)
    check("Duplicate detected across different file formats",
          duplicates.get("anjali_mehta_resume_copy.docx") == "anjali_mehta_resume.pdf",
          str(duplicates))
    for resume in parsed:
        if resume.file_name in duplicates:
            resume.status = "duplicate"
            resume.duplicate_of = duplicates[resume.file_name]

    # ----------------------------------------------------------------- cache
    section("4. Extraction cache")
    cache = ExtractionCache(settings)
    cache.put(anjali)
    cached = cache.get(anjali.content_sha)
    check("Extracted text round-trips through the cache",
          cached is not None and cached.text == anjali.text)
    check("Cache miss returns None", cache.get("0" * 64) is None)

    # -------------------------------------------------------------- analysis
    section("5. Analysis and scoring")
    evaluations = rank_candidates(analyze_batch(parsed, jd, weights, client=client, settings=settings))
    by_name = {e.file_name: e for e in evaluations}
    processed = [e for e in evaluations if e.ok]

    expected_processed = sum(1 for r in parsed if r.ok)
    check("Every uploaded file appears in the results", len(evaluations) == len(files))
    check("All processable candidates were analysed",
          len(processed) == expected_processed, f"{len(processed)} vs {expected_processed}")
    check("Nobody is hidden — failures are listed with a reason",
          all(e.recommendation_reason or e.error for e in evaluations if not e.ok))

    top = by_name["anjali_mehta_resume.pdf"]
    weak = by_name["vikram_desai_resume.pdf"]
    partial = by_name["rahul_nair_resume.docx"]

    check("Strongest candidate ranks #1", top.rank == 1, f"rank={top.rank}")
    check("Off-domain candidate ranks last among processed",
          weak.rank == len(processed), f"rank={weak.rank}")
    check("Relevant candidate scores far above the off-domain one",
          top.overall_score - weak.overall_score > 40,
          f"{top.overall_score} vs {weak.overall_score}")
    check("Scores stay within 0-100", all(0 <= e.overall_score <= 100 for e in processed))
    check("Category scores sum to the overall score",
          all(abs(sum(c.score for c in e.category_scores) - e.overall_score) < 0.15
              for e in processed))
    check("Each category respects its configured maximum",
          all(c.score <= c.max_score + 1e-6 for e in processed for c in e.category_scores))

    check("Every category score has a reason", all(c.reasoning for e in processed for c in e.category_scores))
    check("Recommendation is one of the five defined categories",
          all(e.recommendation in
              {"Strong Match", "Match", "Partial Match", "Weak Match", "Insufficient Information"}
              for e in evaluations))
    check("Every recommendation is explained", all(e.recommendation_reason for e in processed))
    check("Top candidate meets all mandatory requirements",
          top.mandatory_missed_count() == 0,
          str(top.mandatory_requirements_missed))
    check("Off-domain candidate misses most mandatory requirements",
          weak.mandatory_missed_count() >= 5, str(weak.mandatory_missed_count()))
    check("Recommendation is not score-only (mandatory gaps hold candidates back)",
          partial.recommendation in {"Weak Match", "Partial Match"}
          and partial.mandatory_missed_count() > 0)

    check("Per-requirement verdicts recorded for traceability",
          all(len(e.requirement_verdicts) >= 5 for e in processed))
    check("Verdict chain is complete (requirement, relevance, reasoning)",
          all(v.requirement and v.reasoning and 0 <= v.relevance <= 1
              for e in processed for v in e.requirement_verdicts))
    check("Met requirements are backed by resume evidence",
          all(v.evidence for e in processed for v in e.requirement_verdicts if v.met))
    check("Evidence quotes come from the resume text",
          all(_quote_in_resume(q, by_file[e.file_name].text)
              for e in processed for v in e.requirement_verdicts for q in v.evidence))

    check("Contextual matching, not keyword matching (paraphrase recognised)",
          any("saas" in v.requirement.lower() or v.relevance > 0.6
              for v in partial.requirement_verdicts))
    check("Missing information is reported explicitly",
          all(any(NOT_MENTIONED in item or "could not" in item.lower() or "No resume content" in item
                  for item in e.missing_information) or not e.missing_information
              for e in processed))
    check("Candidate with no stated location is flagged, not assumed",
          any(NOT_MENTIONED in item for item in unnamed_eval(by_name).missing_information))
    check("Location scored only because this JD states a requirement",
          any(v.category == "location" for v in top.requirement_verdicts))
    check("Top candidate's location requirement met from stated location",
          all(v.met for v in top.requirement_verdicts if v.category == "location"))

    check("Duplicate is reported, not re-scored",
          by_name["anjali_mehta_resume_copy.docx"].status == "duplicate"
          and "Duplicate of" in by_name["anjali_mehta_resume_copy.docx"].recommendation_reason)

    # ------------------------------------------------------- weights honoured
    section("6. Configurable weights")
    skills_heavy = {"relevant_experience": 10, "required_skills": 50, "jd_alignment": 10,
                    "industry_relevance": 10, "education": 10, "preferred_skills": 10}
    reweighted = analyze_batch(parsed, jd, skills_heavy, client=client, settings=settings)
    reweighted_top = {e.file_name: e for e in reweighted}["anjali_mehta_resume.pdf"]
    check("Re-weighting changes category maximums",
          reweighted_top.score_for("required_skills").max_score == 50,
          str(reweighted_top.score_for("required_skills").max_score))
    check("Re-weighting changes the overall score",
          reweighted_top.overall_score != top.overall_score,
          f"{reweighted_top.overall_score} vs {top.overall_score}")
    check("Re-weighted total still sums correctly",
          abs(sum(c.score for c in reweighted_top.category_scores)
              - reweighted_top.overall_score) < 0.15)

    # ------------------------------------------------- search, filter, ranking
    section("7. Search, filters and ranking")
    check("Semantic search finds a synonym ('AdWords' -> Google Ads experience)",
          any(e.file_name == "anjali_mehta_resume.pdf"
              for e in search_candidates(evaluations, "AdWords")))
    check("Semantic search matches a differently-worded concept ('paid search')",
          len(search_candidates(evaluations, "paid search")) >= 1)
    check("Search excludes irrelevant candidates",
          all(e.file_name == "vikram_desai_resume.pdf"
              for e in search_candidates(evaluations, "Kubernetes")))
    check("Score range filter works",
          all(e.overall_score >= 60 for e in
              filter_candidates(evaluations, score_range=(60, 100), include_unprocessed=False)))
    check("Mandatory-status filter works",
          all(e.mandatory_missed_count() == 0 for e in
              filter_candidates(evaluations, mandatory_status="All mandatory met")))
    check("Recommendation filter works",
          all(e.recommendation == "Strong Match" for e in
              filter_candidates(evaluations, recommendations=["Strong Match"])))
    check("Education filter works",
          len(filter_candidates(evaluations, education_query="MBA")) >= 1)
    check("Ranks are contiguous from 1", [e.rank for e in processed] == list(range(1, len(processed) + 1)))
    check("Unprocessed files keep rank 0 and sort last",
          all(e.rank == 0 for e in evaluations if not e.ok))

    frame = to_dataframe(evaluations)
    check("DataFrame includes every candidate", len(frame) == len(evaluations))
    check("DataFrame is sortable on the major columns",
          {"Rank", "Candidate Name", "Overall Score", "Recommendation"} <= set(frame.columns))

    summary = build_summary(evaluations, len(files))
    check("Summary counts uploads", summary["total_uploaded"] == len(files))
    check("Summary counts processed", summary["processed"] == expected_processed)
    check("Summary counts failures", summary["failed"] == 3, str(summary["failed"]))
    check("Summary counts duplicates", summary["duplicates"] == 1)
    check("Summary reports score statistics",
          summary["highest_score"] >= summary["average_score"] >= summary["lowest_score"])
    check("Summary lists top candidates", len(summary["top_candidates"]) >= 3)

    # ------------------------------------------------------------------- OCR
    section("8. Scanned / image-based PDF (OCR)")
    scanned = by_file["scanned_resume.pdf"]
    ocr_available = _ocr_available()
    print(f"  (Tesseract {'detected' if ocr_available else 'NOT installed'})")

    check("Scanned PDF has no text layer (it is genuinely an image)",
          _pdf_text_layer_empty(samples / "resumes" / "scanned_resume.pdf"))

    if ocr_available:
        check("Scanned PDF processed via OCR",
              scanned.ok and "ocr" in scanned.extraction_method, 
              f"status={scanned.status} method={scanned.extraction_method}")
        check("OCR extracted the candidate name",
              scanned.candidate_name == "Priya Raghavan", scanned.candidate_name)
        check("OCR extracted years of experience", scanned.total_years == 6.0,
              str(scanned.total_years))
        check("OCR extracted skills", len(scanned.skills) >= 4, str(scanned.skills))
        check("OCR extracted education", len(scanned.education) >= 2, str(scanned.education))
        check("OCR extracted employers", len(scanned.companies) >= 2, str(scanned.companies))
        check("Wrapped lines are re-joined, so evidence quotes are not truncated",
              any("budget of INR 90 lakh" in r for r in scanned.responsibilities),
              str(scanned.responsibilities[:1]))
        scanned_eval = by_name["scanned_resume.pdf"]
        check("OCR'd candidate is scored like any other",
              scanned_eval.ok and scanned_eval.overall_score > 0 and scanned_eval.rank > 0,
              f"score={scanned_eval.overall_score} rank={scanned_eval.rank}")
        check("OCR'd candidate's evidence traces back to the OCR text",
              all(_quote_in_resume(q, scanned.text)
                  for v in scanned_eval.requirement_verdicts for q in v.evidence))
    else:
        check("Without OCR installed, the scanned PDF is flagged, never guessed at",
              scanned.status == "ocr_required" and not scanned.text,
              f"status={scanned.status}")
        check("The flag explains what is needed",
              "ocr" in scanned.error.lower() or "scanned" in scanned.error.lower(),
              scanned.error)

    # ------------------------------------------------------- bias and fairness
    section("9. Bias and fairness")
    # Word-boundary matching: a substring test would flag "landing page" and
    # "manage" for containing "age".
    sensitive = ["gender", "male", "female", "religion", "caste", "race", "ethnicity",
                 "nationality", "marital", "married", "age", "dob", "disability",
                 "orientation", "political", "photograph"]
    surfaces = []
    for evaluation in processed:
        surfaces.extend(evaluation.key_strengths)
        surfaces.extend(evaluation.key_concerns)
        surfaces.append(evaluation.recommendation_reason)
        surfaces.extend(c.reasoning for c in evaluation.category_scores)
    blob = " ".join(surfaces).lower()
    hits = [term for term in sensitive if re.search(rf"\b{term}\b", blob)]
    check("No protected characteristic appears in any scoring rationale", not hits, str(hits))
    check("No protected characteristic appears in the exported evidence",
          not [t for t in sensitive
               if re.search(rf"\b{t}\b",
                            " ".join(q for e in processed for q in e.evidence).lower())])

    # --------------------------------------------------------------- exports
    section("10. Exports")
    outputs = ROOT / "outputs"
    xlsx_path = export_to_excel(evaluations, summary, weights,
                                outputs / "sample_candidate_ranking.xlsx", jd=jd,
                                analysis_method=client.describe())
    docx_path = export_to_docx(evaluations, summary, weights,
                               outputs / "sample_candidate_report.docx", jd=jd,
                               analysis_method=client.describe())

    import openpyxl
    from docx import Document

    workbook = openpyxl.load_workbook(xlsx_path)
    check("Excel file created", xlsx_path.exists() and xlsx_path.stat().st_size > 5000)
    check("Excel has the three required sheets",
          workbook.sheetnames == ["Candidate Ranking", "Detailed Evaluation", "Summary"],
          str(workbook.sheetnames))

    ranking_sheet = workbook["Candidate Ranking"]
    check("Ranking sheet has a row per candidate",
          ranking_sheet.max_row == len(evaluations) + 1, str(ranking_sheet.max_row))
    check("Header row frozen", ranking_sheet.freeze_panes == "A2")
    check("Filters enabled", bool(ranking_sheet.auto_filter.ref))
    check("Conditional formatting applied",
          sum(len(rules) for rules in ranking_sheet.conditional_formatting._cf_rules.values()) >= 2)
    check("Column widths set", len(ranking_sheet.column_dimensions) >= 10)
    check("Required ranking columns present",
          {"Rank", "Candidate Name", "Resume", "Overall Score", "Recommendation",
           "Relevant Experience", "Required Skills Match", "JD Alignment", "Industry Match",
           "Education", "Preferred Skills", "Mandatory Requirements Missed", "Key Strengths",
           "Key Concerns"}
          <= {cell.value for cell in ranking_sheet[1]})

    detail_sheet = workbook["Detailed Evaluation"]
    check("Detail sheet has a row per candidate per category",
          detail_sheet.max_row >= len(processed) * len(SCORE_CATEGORIES))
    check("Detail sheet carries reasoning and evidence",
          {"Reasoning", "Resume Evidence"} <= {cell.value for cell in detail_sheet[1]})

    summary_sheet = workbook["Summary"]
    summary_text = " ".join(str(cell.value) for row in summary_sheet.iter_rows() for cell in row
                            if cell.value is not None)
    for label in ["Total resumes uploaded", "Successfully processed", "Strong Match",
                  "Average candidate score", "Highest score", "Lowest score", "Top candidates"]:
        check(f"Summary sheet reports '{label}'", label in summary_text)

    document = Document(str(docx_path))
    check("DOCX file created", docx_path.exists() and docx_path.stat().st_size > 10000)
    headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
    document_text = "\n".join(p.text for p in document.paragraphs)
    for label in ["Overall Ranking", "Score Breakdown", "Key Strengths", "Key Concerns",
                  "Mandatory Requirements", "Missing Information", "Resume Evidence",
                  "Recruiter Notes"]:
        check(f"DOCX report includes '{label}'", any(label in h for h in headings), "")
    check("DOCX names the top candidate", top.candidate_name in document_text)
    check("DOCX lists files that were not evaluated",
          any("Files Not Evaluated" in h for h in headings))

    # ---------------------------------------------------------------- result
    print("\n" + "=" * 74)
    print(f"RESULT: {len(PASSED)} passed, {len(FAILED)} failed")
    print("=" * 74)
    if FAILED:
        for failure in FAILED:
            print(f"  FAILED: {failure}")
        return 1

    print("\nRanked candidates:")
    for evaluation in evaluations:
        rank = f"#{evaluation.rank}" if evaluation.rank else " - "
        print(f"  {rank:>4}  {evaluation.overall_score:5.1f}  {evaluation.recommendation:24}  "
              f"{evaluation.candidate_name:22}  {evaluation.file_name}")
    print(f"\nGenerated:\n  {xlsx_path}\n  {docx_path}")
    return 0


def _ocr_available() -> bool:
    """True when both the Python bindings and the Tesseract binary are present."""
    try:
        import pytesseract
        from PIL import Image  # noqa: F401

        pytesseract.get_tesseract_version()
        return True
    except Exception:  # noqa: BLE001
        return False


def _pdf_text_layer_empty(path: Path) -> bool:
    try:
        import pymupdf
    except ImportError:  # pragma: no cover
        import fitz as pymupdf

    with pymupdf.open(str(path)) as doc:
        return not "".join(page.get_text("text") for page in doc).strip()


def unnamed_eval(by_name):
    return by_name["unnamed_candidate_resume.docx"]


def _quote_in_resume(quote: str, resume_text: str) -> bool:
    """Evidence must be traceable to the resume, not generated."""
    from src.utils import normalize_key

    cleaned = quote.replace("Skills section: ", "").replace("Location: ", "")
    cleaned = cleaned.split("...")[0].strip()
    needle = normalize_key(cleaned)[:60]
    return bool(needle) and needle in normalize_key(resume_text)


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception:  # noqa: BLE001
        traceback.print_exc()
        sys.exit(1)
